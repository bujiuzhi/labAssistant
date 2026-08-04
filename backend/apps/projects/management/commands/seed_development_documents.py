"""按原型初始化可重复执行的项目文档演示数据。"""

from datetime import datetime, timedelta
from io import BytesIO

from django.core.files.base import ContentFile
from django.core.management.base import BaseCommand, CommandError
from django.db import transaction
from django.utils import timezone
from docx import Document
from openpyxl import Workbook
from pptx import Presentation
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.cidfonts import UnicodeCIDFont
from reportlab.pdfgen import canvas

from apps.identity.models import User
from apps.projects.models import (
    Project,
    ProjectDocument,
    ProjectDocumentCategory,
)

CATEGORY_SEQUENCE = [
    ProjectDocumentCategory.PROJECT_PLAN,
    ProjectDocumentCategory.LITERATURE,
    ProjectDocumentCategory.EXPERIMENT_PLAN,
    ProjectDocumentCategory.STAGE_REPORT,
    ProjectDocumentCategory.MEETING_MINUTES,
    ProjectDocumentCategory.OTHER,
]

MIME_TYPES = {
    "docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    "pdf": "application/pdf",
    "xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    "pptx": "application/vnd.openxmlformats-officedocument.presentationml.presentation",
}


def aware(value: str) -> datetime:
    """将原型时间文本转换为当前时区时间。

    Args:
        value: `YYYY-MM-DD HH:MM` 格式时间。

    Returns:
        带时区的日期时间。
    """
    return timezone.make_aware(datetime.strptime(value, "%Y-%m-%d %H:%M"))


def build_real_document(
    *,
    extension: str,
    name: str,
    project: Project,
    category: ProjectDocumentCategory,
) -> bytes:
    """生成可由办公软件实际打开和预览的演示文件。

    Args:
        extension: 目标文件扩展名。
        name: 文档名称。
        project: 关联项目。
        category: 文档分类。

    Returns:
        标准格式文件字节。
    """
    output = BytesIO()
    if extension == "docx":
        document = Document()
        document.add_heading(name.removesuffix(".docx"), level=1)
        document.add_paragraph(f"关联项目：{project.name}")
        document.add_paragraph(f"文档分类：{category.label}")
        document.add_paragraph("本文件用于验证项目文档上传、下载与在线预览的真实链路。")
        table = document.add_table(rows=2, cols=2)
        table.cell(0, 0).text = "项目编号"
        table.cell(0, 1).text = project.project_no
        table.cell(1, 0).text = "项目负责人"
        table.cell(1, 1).text = project.owner.display_name
        document.save(output)
    elif extension == "xlsx":
        workbook = Workbook()
        worksheet = workbook.active
        worksheet.title = "第三轮配方筛选"
        worksheet.append(["样品编号", "原料", "批次", "用量/g", "结果"])
        worksheet.append(["A-01", "聚合物基体", "B-202607", 60, "待检测"])
        worksheet.append(["A-02", "改性剂", "M-202607", 5, "合格"])
        worksheet.freeze_panes = "A2"
        workbook.save(output)
    elif extension == "pptx":
        presentation = Presentation()
        slide = presentation.slides.add_slide(presentation.slide_layouts[1])
        slide.shapes.title.text = name.removesuffix(".pptx")
        slide.placeholders[1].text = (
            f"项目：{project.name}\n"
            f"分类：{category.label}\n"
            "用于验证演示文稿在线预览与下载。"
        )
        presentation.save(output)
    elif extension == "pdf":
        pdfmetrics.registerFont(UnicodeCIDFont("STSong-Light"))
        pdf = canvas.Canvas(output)
        pdf.setFont("STSong-Light", 18)
        pdf.drawString(72, 780, name.removesuffix(".pdf"))
        pdf.setFont("STSong-Light", 12)
        pdf.drawString(72, 744, f"关联项目：{project.name}")
        pdf.drawString(72, 720, f"文档分类：{category.label}")
        pdf.drawString(72, 696, "本文件用于验证 PDF 原文件在线预览。")
        pdf.save()
    else:
        raise ValueError(f"不支持生成 {extension} 演示文件")
    return output.getvalue()


class Command(BaseCommand):
    """初始化项目文档演示数据。"""

    help = "初始化项目详情文档资料演示数据"

    @transaction.atomic
    def handle(self, *args, **options) -> None:
        """执行初始化。

        Args:
            *args: 未使用的位置参数。
            **options: 未使用的命令参数。

        Raises:
            CommandError: 项目或开发账号尚未初始化。
        """
        actor = User.objects.filter(username="admin", is_active=True).first()
        if not actor:
            raise CommandError("请先执行 bootstrap_development 初始化开发账号")
        projects = list(
            Project.objects.filter(organization=actor.organization).order_by("project_no")
        )
        if not projects:
            raise CommandError("请先执行 seed_development_projects 初始化项目")

        base_time = aware("2026-07-20 10:24")
        created_count = 0
        updated_count = 0
        for project in projects:
            target_count = max(project.document_count, 4)
            core_specs = [
                (
                    "项目实施方案V3.docx",
                    ProjectDocumentCategory.PROJECT_PLAN,
                    "V3.0",
                    "docx",
                ),
                (
                    f"{project.name[:18]}研究综述.pdf",
                    ProjectDocumentCategory.LITERATURE,
                    "V1.0",
                    "pdf",
                ),
                (
                    "第三轮配方筛选记录.xlsx",
                    ProjectDocumentCategory.EXPERIMENT_PLAN,
                    "V2.1",
                    "xlsx",
                ),
                (
                    "项目中期汇报-202607.pptx",
                    ProjectDocumentCategory.STAGE_REPORT,
                    "V1.2",
                    "pptx",
                ),
            ]
            specs = list(core_specs)
            for index in range(max(0, target_count - len(core_specs))):
                category = CATEGORY_SEQUENCE[index % len(CATEGORY_SEQUENCE)]
                extension = ["docx", "pdf", "xlsx", "pptx"][index % 4]
                specs.append(
                    (
                        f"{category.label}归档资料-{index + 1:02d}.{extension}",
                        category,
                        f"V1.{index + 1}",
                        extension,
                    )
                )

            for index, (
                name,
                category,
                version_label,
                extension,
            ) in enumerate(specs):
                document = ProjectDocument.objects.filter(
                    project=project,
                    name=name,
                ).first()
                payload = build_real_document(
                    extension=extension,
                    name=name,
                    project=project,
                    category=category,
                )
                if document:
                    document.file.delete(save=False)
                    updated_count += 1
                else:
                    document = ProjectDocument(
                        organization=project.organization,
                        project=project,
                        uploaded_by=project.owner,
                    )
                    created_count += 1
                document.name = name
                document.extension = extension
                document.mime_type = MIME_TYPES[extension]
                document.file_size = len(payload)
                document.category = category
                document.version_label = version_label
                document.updated_by = project.owner
                document.file.save(name, ContentFile(payload), save=False)
                document.save()
                updated_at = base_time - timedelta(hours=index * 7)
                ProjectDocument.objects.filter(id=document.id).update(
                    created_at=updated_at,
                    updated_at=updated_at,
                )
            project.document_count = project.documents.count()
            project.save(update_fields=["document_count", "updated_at"])

        self.stdout.write(
            self.style.SUCCESS(
                f"项目文档初始化完成，新增 {created_count} 个，更新 {updated_count} 个"
            )
        )
